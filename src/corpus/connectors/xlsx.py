"""Excel `.xlsx` connector via `openpyxl` (MIT).

One workbook becomes one SourceDocument. Each sheet contributes its name as a
heading followed by its non-empty rows, cells tab-joined.

Spreadsheets retrieve poorly whatever you do — the goal here is findability, not
fidelity. Formulas are read as their cached values (`data_only=True`), so a
workbook never opened by Excel may have empty formula cells.

Install: `pip install corpus-rag[xlsx]` or `uv add openpyxl`.
"""

from __future__ import annotations

import logging
import os
from collections.abc import Iterable
from pathlib import Path

from corpus.connectors.discovery import discover_files
from corpus.types import SourceDocument
from corpus.util.dedup import NearDuplicates
from corpus.util.ooxml import permanent_read_failure_reason

logger = logging.getLogger(__name__)


class XlsxConnector:
    def __init__(
        self,
        source_type: str,
        path: Path | str,
        glob: str = "**/*.xlsx",
    ):
        self.source_type = source_type
        self._root = Path(os.path.expanduser(str(path))).resolve()
        self._glob = glob
        self.failed_files = 0
        self.skipped_files = 0

    def load(self) -> Iterable[SourceDocument]:
        # Per-file read failures are counted, not just logged: a skipped file
        # yields no document, so its chunk ids vanish from `seen_ids` and the
        # ingester's orphan sweep would delete already-indexed content. The
        # ingester reads this counter and suppresses pruning when it is
        # non-zero. Reset per run so a reused instance cannot suppress pruning
        # forever on the strength of an old failure.
        self.failed_files = 0
        self.skipped_files = 0
        import openpyxl

        if not self._root.is_dir():
            raise FileNotFoundError(
                f"Xlsx source '{self.source_type}': directory not found: {self._root}"
            )

        dupes = NearDuplicates(self.source_type)
        for path in discover_files(self._root, self._glob):
            # Permanent conditions are skipped, not failed: they recur
            # identically on every run, so counting them as failures
            # would suppress this source's orphan pruning forever.
            reason = permanent_read_failure_reason(path)
            if reason is not None:
                logger.info(
                    "%s: skipping '%s' — %s", self.source_type, path.name, reason
                )
                self.skipped_files += 1
                continue
            try:
                wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
            except Exception as e:
                logger.warning("Xlsx source '%s': cannot open %s: %s", self.source_type, path, e)
                self.failed_files += 1
                continue

            parts: list[str] = []
            try:
                for ws in wb.worksheets:
                    rows: list[str] = []
                    for row in ws.iter_rows(values_only=True):
                        cells = ["" if v is None else str(v).strip() for v in row]
                        if any(cells):
                            rows.append("\t".join(cells).rstrip("\t"))
                    if rows:
                        parts.append(f"{ws.title}\n" + "\n".join(rows))
            except Exception as e:
                logger.warning("Xlsx source '%s': error reading %s: %s", self.source_type, path, e)
                self.failed_files += 1
                continue
            finally:
                wb.close()

            body = "\n\n".join(parts).strip()
            if not body:
                logger.info("%s: skipping '%s' — no cell content", self.source_type, path.name)
                continue

            source_key = str(path.relative_to(self._root))
            if dupes.seen_before(body, source_key):
                continue

            yield SourceDocument(
                source_type=self.source_type,
                source_key=source_key,
                title=path.stem,
                url=None,
                raw={"body": body, "path": str(path)},
            )
