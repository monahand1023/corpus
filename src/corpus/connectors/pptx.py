"""PowerPoint `.pptx` connector via `python-pptx` (MIT).

One deck becomes one SourceDocument. Each slide is its own markdown `##`
section — `title\n\nbody\n\n**Speaker notes:** ...` — so the shared markdown
chunker (see `markdown.py`) splits per slide wherever it reasonably can, and a
retrieval hit carries enough of the slide's own heading to say which slide it
came from. Small adjacent slides still get packed together by the chunker's
own coalescing, the same as xlsx's per-sheet sections and docx's paragraphs +
tables — this connector doesn't special-case that.

Speaker notes are extracted alongside slide text, not left out. A deck's
notes routinely carry the actual narrative — what the presenter meant to say
— while the slide itself is bullet fragments; without notes, a chunk built
from slide text alone is often close to unsearchable ("Q3", "62%", "Next
steps"). Both `has_text_frame` shapes and tables are read; grouped shapes are
walked recursively (`MSO_SHAPE_TYPE.GROUP`) since real-world decks group
logos/labels/etc. and text living only inside a group would otherwise vanish
silently.

Extraction order per slide: the title placeholder (if any) becomes the
section heading and is excluded from the body text extraction to avoid
appearing twice; every other shape's text (in shape order) becomes the body;
speaker notes are appended last. A slide with a title but no body/notes still
gets a one-line section (title-only "section divider" slides are common and
searchable in their own right); a slide with NOTHING extractable (an
image-only slide with no notes) is dropped from the body entirely. If every
slide in the deck is like that, the whole file is skipped — same treatment as
an empty docx.

Legacy `.doc`-style problem for PowerPoint: the old binary `.ppt` format
(OLE2/CFBF container) is not something `python-pptx` can ever read — it only
reads the OPC/zip package that `.pptx` is. The default glob here is
`**/*.pptx` only, so a `.ppt` file is invisible by default (exactly like
`docx.py` and legacy `.doc`) and this case mostly matters when a glob is
widened or a file was misnamed. `python-pptx` itself can't tell "this is an
old-format binary" apart from "this is corrupt garbage" — both raise the same
`PackageNotFoundError` (verified: a `.ppt`-magic file, random garbage, a
truncated zip, and an empty file all raise the identical exception with the
identical message). So the OLE2 signature (`D0 CF 11 E0 A1 B1 1A E1`, the
standard Compound File Binary magic that every `.ppt`/`.doc`/`.xls` starts
with) is checked directly, BEFORE ever calling `Presentation()`: a match means
this file will never parse no matter how many times ingestion retries, so
it's counted in `skipped_files` (permanent), not `failed_files` (possibly
transient — corrupted mid-write, momentarily locked, a genuinely malformed
`.pptx`). See docx.py's/xlsx.py's own docstrings and `ingester.py`'s
`_reported_failures`/`_reported_skips` for why that distinction matters to
orphan pruning.

Install: `pip install corpus-rag[pptx]` or `uv add python-pptx`.
"""

from __future__ import annotations

import logging
import os
from collections.abc import Iterable
from pathlib import Path

from corpus.connectors.discovery import discover_files
from corpus.types import SourceDocument
from corpus.util.dedup import fingerprint

logger = logging.getLogger(__name__)

# Standard Compound File Binary (OLE2) magic — every legacy `.ppt`/`.doc`/
# `.xls` starts with these 8 bytes. See the module docstring: this is the
# only reliable way to tell "old binary format, will never parse" apart from
# "corrupt/garbage `.pptx`", since python-pptx raises the identical
# `PackageNotFoundError` for both.
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _extract_slide_text(slide: object) -> tuple[str | None, list[str]]:
    """Return `(title_text_or_None, body_lines)` for one slide.

    Walks every shape on the slide: text frames become a body line, tables
    become tab-joined rows, and a `GROUP` shape is recursed into (its own
    children get the same treatment) rather than skipped. The title
    placeholder, if present, is excluded from `body_lines` — its text is
    returned separately and becomes the slide's section heading instead of
    being duplicated in the body.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    title_shape = slide.shapes.title  # type: ignore[attr-defined]
    title_id = None
    title_text: str | None = None
    if title_shape is not None:
        title_id = getattr(title_shape, "shape_id", None)
        if getattr(title_shape, "has_text_frame", False):
            t = title_shape.text_frame.text.strip()
            title_text = t or None

    lines: list[str] = []

    def walk(shapes: object) -> None:
        for shape in shapes:  # type: ignore[attr-defined]
            if title_id is not None and getattr(shape, "shape_id", None) == title_id:
                continue
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes)
                continue
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        lines.append("\t".join(cells))

    walk(slide.shapes)  # type: ignore[attr-defined]
    return title_text, lines


def _extract_notes(slide: object) -> str:
    if not getattr(slide, "has_notes_slide", False):
        return ""
    return str(slide.notes_slide.notes_text_frame.text).strip()  # type: ignore[attr-defined]


def _slide_section(index: int, title: str | None, body_lines: list[str], notes: str) -> str | None:
    """One markdown `##` section for a slide, or None if the slide has
    nothing extractable at all (no title, no body text, no notes — e.g. a
    purely decorative or image-only slide)."""
    if title is None and not body_lines and not notes:
        return None
    heading = f"## Slide {index}: {title}" if title else f"## Slide {index}"
    parts = [heading]
    if body_lines:
        parts.append("\n\n".join(body_lines))
    if notes:
        parts.append(f"**Speaker notes:** {notes}")
    return "\n\n".join(parts)


class PptxConnector:
    def __init__(
        self,
        source_type: str,
        path: Path | str,
        glob: str = "**/*.pptx",
    ):
        self.source_type = source_type
        self._root = Path(os.path.expanduser(str(path))).resolve()
        self._glob = glob
        self.failed_files = 0
        self.skipped_files = 0

    def load(self) -> Iterable[SourceDocument]:
        # Per-file failures are counted, not just logged — see the identical
        # note in docx.py/xlsx.py: the ingester reads these to gate orphan
        # pruning. Reset per run so a reused instance can't suppress pruning
        # forever on the strength of an old run's failures.
        self.failed_files = 0
        self.skipped_files = 0
        from pptx import Presentation

        if not self._root.is_dir():
            raise FileNotFoundError(
                f"Pptx source '{self.source_type}': directory not found: {self._root}"
            )

        seen: dict[str, str] = {}
        for path in discover_files(self._root, self._glob):
            try:
                with path.open("rb") as f:
                    header = f.read(len(_OLE2_MAGIC))
            except OSError as e:
                logger.warning(
                    "Pptx source '%s': cannot read %s: %s", self.source_type, path, e
                )
                self.failed_files += 1
                continue

            if header.startswith(_OLE2_MAGIC):
                logger.info(
                    "%s: skipping '%s' — legacy binary .ppt (OLE2) format; "
                    "python-pptx cannot read this and never will",
                    self.source_type,
                    path.name,
                )
                self.skipped_files += 1
                continue

            try:
                presentation = Presentation(str(path))
                # python-pptx parses lazily too: materialize everything we
                # need inside this guard, same rationale as docx.py's
                # equivalent note — a malformed package can construct fine
                # and only raise once a slide/shape is actually touched.
                sections: list[str] = []
                for i, slide in enumerate(presentation.slides, start=1):
                    title, lines = _extract_slide_text(slide)
                    notes = _extract_notes(slide)
                    section = _slide_section(i, title, lines, notes)
                    if section:
                        sections.append(section)
            except Exception as e:  # python-pptx raises many types for bad files
                logger.warning(
                    "Pptx source '%s': cannot open %s: %s", self.source_type, path, e
                )
                self.failed_files += 1
                continue

            body = "\n\n".join(sections).strip()
            if not body:
                logger.info(
                    "%s: skipping '%s' — no extractable text (image-only slides?)",
                    self.source_type,
                    path.name,
                )
                continue

            source_key = str(path.relative_to(self._root))
            fp = fingerprint(body)
            if fp in seen:
                logger.info(
                    "%s: skipping near-duplicate '%s' (matches '%s')",
                    self.source_type,
                    source_key,
                    seen[fp],
                )
                continue
            seen[fp] = source_key

            title = path.stem
            try:
                meta_title = presentation.core_properties.title
                if meta_title and str(meta_title).strip():
                    title = str(meta_title).strip()
            except Exception:
                pass

            yield SourceDocument(
                source_type=self.source_type,
                source_key=source_key,
                title=title,
                url=None,
                raw={"body": body, "path": str(path)},
            )
