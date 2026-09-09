from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from corpus.connectors.pptx import _OLE2_MAGIC, PptxConnector


def _save_deck(path: Path, build) -> None:  # type: ignore[no-untyped-def]
    prs = Presentation()
    build(prs)
    prs.save(path)


def test_loads_slide_text_and_speaker_notes(tmp_path: Path) -> None:
    def build(prs):  # type: ignore[no-untyped-def]
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Quarterly Overview"
        slide.placeholders[1].text_frame.text = "Revenue grew 12%"
        slide.notes_slide.notes_text_frame.text = "Mention the one-time refund in Q2."
        prs.core_properties.title = "Q3 Board Deck"

    _save_deck(tmp_path / "deck.pptx", build)
    docs = list(PptxConnector(source_type="decks", path=tmp_path).load())

    assert len(docs) == 1
    body = docs[0].raw["body"]
    assert docs[0].title == "Q3 Board Deck"
    assert "## Slide 1: Quarterly Overview" in body
    assert "Revenue grew 12%" in body
    assert "**Speaker notes:** Mention the one-time refund in Q2." in body


def test_title_falls_back_to_filename_stem(tmp_path: Path) -> None:
    def build(prs):  # type: ignore[no-untyped-def]
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Untitled"
        slide.placeholders[1].text_frame.text = "Body text."

    _save_deck(tmp_path / "untitled-deck.pptx", build)
    docs = list(PptxConnector(source_type="decks", path=tmp_path).load())
    assert docs[0].title == "untitled-deck"


def test_table_cells_are_extracted(tmp_path: Path) -> None:
    def build(prs):  # type: ignore[no-untyped-def]
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Numbers"
        table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
        table = table_shape.table
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Widget count"
        table.cell(1, 1).text = "42"

    _save_deck(tmp_path / "numbers.pptx", build)
    docs = list(PptxConnector(source_type="decks", path=tmp_path).load())
    body = docs[0].raw["body"]
    assert "Metric\tValue" in body
    assert "Widget count\t42" in body


def test_grouped_shapes_are_walked_recursively(tmp_path: Path) -> None:
    def build(prs):  # type: ignore[no-untyped-def]
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb1 = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1))
        tb1.text_frame.text = "Grouped label A"
        tb2 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        tb2.text_frame.text = "Grouped label B"
        slide.shapes.add_group_shape([tb1, tb2])

    _save_deck(tmp_path / "grouped.pptx", build)
    docs = list(PptxConnector(source_type="decks", path=tmp_path).load())
    body = docs[0].raw["body"]
    assert "Grouped label A" in body
    assert "Grouped label B" in body


def test_title_only_slide_is_still_indexed(tmp_path: Path) -> None:
    """A section-divider slide (title, no body, no notes) carries real
    structural content and must not be dropped."""

    def build(prs):  # type: ignore[no-untyped-def]
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Part Two: Financials"

    _save_deck(tmp_path / "divider.pptx", build)
    docs = list(PptxConnector(source_type="decks", path=tmp_path).load())
    assert len(docs) == 1
    assert "## Slide 1: Part Two: Financials" in docs[0].raw["body"]


def test_skips_image_only_deck(tmp_path: Path) -> None:
    def build(prs):  # type: ignore[no-untyped-def]
        prs.slides.add_slide(prs.slide_layouts[6])  # blank layout, no shapes at all

    _save_deck(tmp_path / "images_only.pptx", build)
    docs = list(PptxConnector(source_type="decks", path=tmp_path).load())
    assert docs == []


def test_legacy_ppt_binary_is_skipped_not_failed(tmp_path: Path) -> None:
    """A `.ppt`-format (OLE2) file will never parse, no matter how many
    times ingestion retries -- counted in skipped_files, not failed_files."""
    (tmp_path / "renamed_legacy.pptx").write_bytes(_OLE2_MAGIC + b"\x00" * 64)
    conn = PptxConnector(source_type="decks", path=tmp_path)
    docs = list(conn.load())
    assert docs == []
    assert conn.skipped_files == 1
    assert conn.failed_files == 0


def test_corrupt_file_is_failed_not_skipped(tmp_path: Path) -> None:
    (tmp_path / "corrupt.pptx").write_bytes(b"this is not a pptx package at all")
    conn = PptxConnector(source_type="decks", path=tmp_path)
    docs = list(conn.load())
    assert docs == []
    assert conn.failed_files == 1
    assert conn.skipped_files == 0


def test_missing_dir_raises() -> None:
    with pytest.raises(FileNotFoundError):
        list(PptxConnector(source_type="decks", path="/nonexistent").load())


def test_dedupes_identical_decks(tmp_path: Path) -> None:
    def build(prs):  # type: ignore[no-untyped-def]
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Same Deck"
        slide.placeholders[1].text_frame.text = "Identical content."

    _save_deck(tmp_path / "a.pptx", build)
    _save_deck(tmp_path / "b.pptx", build)
    docs = list(PptxConnector(source_type="decks", path=tmp_path).load())
    assert len(docs) == 1
